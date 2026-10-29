#!/usr/bin/env python3
"""Freelance DevOps / AWS work finder - projects and part-time gigs, worldwide.

    python3 freelance_leads/freelance_leads.py            find work, draft emails, write the Excel (sends nothing)
    python3 freelance_leads/freelance_leads.py --send     ...and email every new HR / client contact found

Each run looks for DevOps / AWS / cloud work you can do on the side (2-3 hours a day or weekends),
from clients anywhere - US, UK, Europe, Australia, India, ...:
  1. freelance marketplaces: Freelancer.com (public API), PeoplePerHour (project pages), and - through
     Google, past week - Upwork, Guru, Contra, Truelancer, freelancermap and Reddit r/forhire;
     plus LinkedIn posts, part-time / contract roles on LinkedIn and Indeed (US, UK, EU, AU, CA, IN)
     and the Remotive / Jobicy / Himalayas feeds;
  2. keeps only real, current, side-work friendly leads: a single project or post (not a listing page),
     posted in the last 10 days, still open, remote, not full-time, for 2-4 years of experience,
     a sane budget, no scam wording, and a DevOps match with your resume - every kept lead records
     where it was verified (API / page data / search index date);
  3. emails the leads that publish a contact address (short HTML note + resume) - marketplace
     projects are applied to on the platform (link in the sheet);
  4. writes runs/<stamp>/devops_leads_<stamp>.xlsx (SR, Company Name, Position, Works / Description,
     Location, Salary / Rate, Contract Type, Contact / HR Email, Job Role, Status) and emails it to you.
     If fewer than `min_leads` (15) survive, it searches more skill terms until it has them or the
     search budget is spent.

Keys come only from variables that already exist: EMAIL_ADDRESS, EMAIL_PASSWORD (SMTP_HOST/
SMTP_PORT optional), SERPER_KEY / SERPAPI_KEY / SCRAPERAPI_KEY / SCRAPINGDOG_KEY / SEARCHAPI_KEY,
GEMINI_API_KEY / GROQ_API_KEY / OPENROUTER_API_KEY, REPORT_EMAIL (optional). Your details and
resume come from job_autopilot/profile.yaml and bulk_mail_sender/Saurabh_Agrawal_2026.pdf.
"""

from __future__ import annotations

import os
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
APP = REPO / "job_autopilot"
VENV = APP / ".venv"

if os.path.realpath(sys.prefix) != os.path.realpath(VENV) and (VENV / "bin" / "python").exists():
    os.execv(str(VENV / "bin" / "python"), [str(VENV / "bin" / "python"), str(Path(__file__).resolve()), *sys.argv[1:]])
sys.path.insert(0, str(APP))

import argparse  # noqa: E402
import html as htmllib  # noqa: E402
import json  # noqa: E402
import logging  # noqa: E402
import random  # noqa: E402
import re  # noqa: E402
import smtplib  # noqa: E402
import ssl  # noqa: E402
import threading  # noqa: E402
import time  # noqa: E402
from collections import Counter  # noqa: E402
from dataclasses import dataclass  # noqa: E402
from datetime import datetime, timedelta, timezone  # noqa: E402
from email.message import EmailMessage  # noqa: E402
from email.utils import formatdate  # noqa: E402
from types import SimpleNamespace  # noqa: E402
from urllib.parse import urlsplit  # noqa: E402

import requests  # noqa: E402
import yaml  # noqa: E402
from requests.adapters import HTTPAdapter  # noqa: E402
from urllib3.util.retry import Retry  # noqa: E402

from autopilot.config import load_env, load_profile, load_settings  # noqa: E402
from autopilot.discovery import ai_search, boards, google_search as gs  # noqa: E402
from autopilot.discovery.common import USER_AGENT, mail_domain_ok, to_datetime  # noqa: E402
from autopilot.models import Job, canonical_url  # noqa: E402
from autopilot.text import CORE_TERMS, EMAIL_RE, display, find_terms, html_to_text, implied, pick_emails  # noqa: E402

IST = timezone(timedelta(hours=5, minutes=30), "IST")
COLUMNS = ["SR", "Company Name", "Position", "Works / Description", "Location", "Salary / Rate",
           "Contract Type", "Contact / HR Email", "Job Role", "Status"]
SENT, DRAFTED, REVIEW = "Sent", "Email Drafted", "Pending Review"
log = logging.getLogger("freelance")


# --------------------------------------------------------------------------- #
# Logging: in GitHub Actions (public logs) email addresses are masked
# --------------------------------------------------------------------------- #
class _MaskFormatter(logging.Formatter):
    def format(self, record):
        return EMAIL_RE.sub("<email>", super().format(record))


def setup_logging(run_dir: Path) -> None:
    log.setLevel(logging.INFO)
    fmt = logging.Formatter("%(asctime)s %(levelname)s %(message)s", "%H:%M:%S")
    console = logging.StreamHandler(sys.stdout)
    console.setFormatter(_MaskFormatter(fmt._fmt, fmt.datefmt) if os.getenv("CI") else fmt)
    file = logging.FileHandler(run_dir / "run.log", encoding="utf-8")   # full log: goes to your inbox only
    file.setFormatter(fmt)
    log.handlers[:] = [console, file]


# --------------------------------------------------------------------------- #
# Polite HTTP: retries + at most one request per host every 1.5 s
# --------------------------------------------------------------------------- #
class ThrottledSession(requests.Session):
    def __init__(self, gap: float = 1.5):
        super().__init__()
        self.headers.update({"User-Agent": USER_AGENT, "Accept-Language": "en-US,en;q=0.9"})
        retry = Retry(total=2, backoff_factor=1.5, status_forcelist=(500, 502, 503, 504), allowed_methods=("GET",))
        self.mount("https://", HTTPAdapter(max_retries=retry))
        self.gap, self._last, self._lock = gap, {}, threading.Lock()

    def request(self, method, url, *args, **kwargs):
        host = urlsplit(url).netloc
        with self._lock:
            wait = self._last.get(host, 0) + self.gap - time.time()
            self._last[host] = max(time.time(), self._last.get(host, 0) + self.gap)
        if wait > 0:
            time.sleep(wait)
        kwargs["timeout"] = min(float(kwargs.get("timeout") or 25), 30.0)   # one slow provider can't stall a run
        return super().request(method, url, *args, **kwargs)


# --------------------------------------------------------------------------- #
# What kind of work, what it pays, does it fit 2-3 hours a day
# --------------------------------------------------------------------------- #
_TYPE_RULES = [
    ("Freelance", re.compile(r"\bfreelanc\w*|\bgig\b|independent contractor|\[hiring\]", re.I)),
    ("Part-Time", re.compile(r"part[\s-]*time|\b\d{1,2}\s*(?:-\s*\d{1,2}\s*)?(?:hrs?|hours)\s*(?:/|per|a)\s*week", re.I)),
    ("Project", re.compile(r"project[\s-]*based|one[\s-]*time project|short[\s-]*term (?:project|engagement)|"
                           r"fixed[\s-]*price|per project|milestone", re.I)),
    ("Contract", re.compile(r"\bc2h\b|contract[\s-]*to[\s-]*hire|\bc2c\b|corp[\s-]*to[\s-]*corp|\bcontractual\b|"
                            r"(?<!employment )(?<!smart )\bcontract\b(?!\s*(?:law|management|negotiat|lifecycle|"
                            r"review|drafting|manufactur))|\btemp(?:orary)? role\b", re.I)),
]
_FEED_TYPE = {"contract": "Contract", "freelance": "Freelance", "part_time": "Part-Time", "part-time": "Part-Time",
              "parttime": "Part-Time", "contractor": "Contract", "temporary": "Contract", "hourly": "Freelance - Hourly",
              "fixed": "Freelance - Fixed price"}


def contract_type(job: Job, declared: str = "") -> str:
    """Freelance (hourly / fixed price) / Part-Time / Project / Contract, or '' for a regular job."""
    if declared:
        return _FEED_TYPE.get(declared.lower().replace(" ", "_"), _FEED_TYPE.get(declared.lower(), declared))
    text = f"{job.title}\n{job.description[:2500]}"
    if re.search(r"\bhourly\b|/\s*hr\b|per hour", text, re.I) and job.source in PLATFORM_SOURCES:
        return "Freelance - Hourly"
    if re.search(r"fixed[\s-]*price|\bbudget\b", text, re.I) and job.source in PLATFORM_SOURCES:
        return "Freelance - Fixed price"
    for part in (job.title, job.description[:2500]):
        for label, rx in _TYPE_RULES:
            if rx.search(part or ""):
                return label
    return "Freelance" if job.source in PLATFORM_SOURCES else ""


_MONEY = r"(?:₹|rs\.?|inr|\$|usd|€|eur|£|gbp|aed|sgd|aud|cad)\s?\d[\d,]*(?:\.\d+)?\s*(?:k|l|lakhs?|lpa)?"
_RATE = re.compile(rf"{_MONEY}(?:\s*(?:-|–|to)\s*(?:{_MONEY}|\d[\d,]*(?:\.\d+)?\s*(?:k|l|lakhs?)?))?"
                   rf"\s*(?:/|per|an?)\s*(?:hr|hour|h|day|week|month|mo|annum|year|yr)\b|"
                   rf"{_MONEY}\s*(?:-|–|to)\s*{_MONEY}(?:\s*(?:hourly|fixed[\s-]*price))?|"
                   rf"budget[:\s-]*(?:of\s*)?{_MONEY}|"
                   rf"\d+(?:\.\d+)?\s*(?:-|–|to)\s*\d+(?:\.\d+)?\s*(?:lpa|lakhs? per annum)", re.I)


def rate_of(job: Job) -> str:
    if job.salary_min or job.salary_max:
        lo, hi = job.salary_min, job.salary_max
        amount = f"{lo:,.0f}-{hi:,.0f}" if lo and hi and lo != hi else f"{(lo or hi):,.0f}"
        per = {"hourly": "/hr", "fixed": "fixed price"}.get(job.salary_interval, f"/ {job.salary_interval or 'period n/a'}")
        return f"{job.salary_currency or ''} {amount} {per}".strip()
    m = _RATE.search(f"{job.title}\n{job.description[:6000]}")
    return re.sub(r"\s+", " ", m.group(0)).strip() if m else "Not stated"


def usd_budget(job: Job) -> float | None:
    """The top of the budget in USD (per hour for hourly work), when the source states one."""
    from autopilot.screening import USD_RATE

    top = job.salary_max or job.salary_min
    rate = USD_RATE.get((job.salary_currency or "").upper())
    if top and rate:
        return float(top) * rate
    m = re.search(r"\$\s?(\d[\d,]*(?:\.\d+)?)(?:\s*(?:-|–|to)\s*\$\s?(\d[\d,]*(?:\.\d+)?))?", job.title + " " + job.description[:3000])
    if m:
        return float((m.group(2) or m.group(1)).replace(",", ""))
    return None


_FULL_TIME = re.compile(r"\bfull[\s-]*time\b(?![\s-]*equivalent)|\b40\s*\+?\s*(?:hrs?|hours)\s*(?:/|per|a)\s*week|"
                        r"more than 30 hrs|30\+\s*hrs|\b8\s*hours?\s*(?:a|per)\s*day|\bmonday\s*(?:-|to)\s*friday\b|"
                        r"\bin[\s-]*office\b|\bon[\s-]*site\b|\bwork from office\b", re.I)
_SIDE_TIME = re.compile(r"part[\s-]*time|less than 30 hrs|flexible (?:hours|schedule|time)|few hours|weekends?|"
                        r"\b(?:[1-9]|1\d|2\d)\s*(?:-\s*\d{1,2}\s*)?(?:hrs?|hours)\s*(?:/|per|a)\s*week|as needed|"
                        r"ad[\s-]*hoc|one[\s-]*time|short[\s-]*term|fixed[\s-]*price|milestone|hourly|per hour|"
                        r"\bproject\b|\btask\b|\bgig\b|freelanc", re.I)
_SENIOR = re.compile(r"\b(senior|sr\.?|lead|principal|staff|architect|head|manager|director|vp|chief)\b", re.I)


def side_time_reason(lead: "Lead") -> str:
    """'' when the work fits 2-3 hours a day / weekends; a reason otherwise."""
    job = lead.job
    text = f"{job.title}\n{job.description[:4000]}"
    if job.source in PLATFORM_SOURCES:
        return "full-time commitment" if _FULL_TIME.search(text) and not _SIDE_TIME.search(text) else ""
    if re.search(r"\bfull[\s-]*time\b", text, re.I) and not re.search(r"part[\s-]*time", text, re.I):
        return "full-time contract (needs 40 h/week)"
    if lead.kind == "Part-Time" and lead.declared == "part_time":
        return ""                                     # the posting itself says part-time
    if _SIDE_TIME.search(text):
        return ""
    return "full-time contract (needs 40 h/week)"


_FAMILY = [("SRE / Reliability", r"\bsre\b|reliability"), ("MLOps", r"mlops"), ("DevSecOps", r"devsecops|security"),
           ("Platform / Kubernetes", r"platform|kubernetes|k8s|helm|eks\b|aks\b|gke\b"),
           ("CI/CD pipelines", r"ci/cd|cicd|pipeline|jenkins|github actions|gitlab"),
           ("Infrastructure as Code", r"terraform|iac|ansible|cloudformation|pulumi"),
           ("Cloud Infrastructure", r"cloud|aws|azure|gcp|lambda|serverless|ec2|s3\b")]


def job_role(job: Job) -> str:
    family = next((name for name, rx in _FAMILY if re.search(rx, job.title, re.I)), "DevOps / Cloud")
    skills = display(job.matched[:3])
    return f"{family} - {', '.join(skills)}" if skills else family


def location_label(job: Job) -> str:
    if job.source in PLATFORM_SOURCES:
        return f"Remote - worldwide ({job.source})"
    if job.is_remote:
        return "Remote (Global)" if re.search(r"worldwide|anywhere|global", job.location or "", re.I) else \
            f"Remote ({job.location.strip()})" if job.location and job.location.lower() != "remote" else "Remote"
    return job.location or "Not stated"


def brief(text: str, limit: int = 220) -> str:
    text = re.sub(r"\s+", " ", html_to_text(text or "")).strip()
    sentences = re.split(r"(?<=[.!?])\s+", text)
    out = ""
    for s in sentences:
        if len(out) + len(s) > limit:
            break
        out += (" " if out else "") + s
    return out or text[:limit]


# --------------------------------------------------------------------------- #
# Leads
# --------------------------------------------------------------------------- #
PLATFORM_SOURCES = {"Upwork", "Freelancer.com", "PeoplePerHour", "Guru", "Contra", "Truelancer", "freelancermap",
                    "Reddit r/forhire"}
# One project per URL (a listing / search page is not a lead).
PLATFORMS = {
    "upwork.com": ("Upwork", re.compile(r"/freelance-jobs/apply/", re.I)),
    "guru.com": ("Guru", re.compile(r"/jobs?/.+/\d{5,}|/d/jobs/.+/\d{5,}", re.I)),
    "contra.com": ("Contra", re.compile(r"/opportunit", re.I)),
    "truelancer.com": ("Truelancer", re.compile(r"/freelance-project/", re.I)),
    "freelancermap.com": ("freelancermap", re.compile(r"/projects?/.*\d{5,}|/project/", re.I)),
    "freelancermap.de": ("freelancermap", re.compile(r"/projekt", re.I)),
    "reddit.com": ("Reddit r/forhire", re.compile(r"/comments/", re.I)),
}
_SUFFIX = re.compile(r"\s*(?:[-|–:]\s*)?(?:freelance job in .*|upwork|guru|contra|truelancer(?:\.com)?|freelancermap|"
                     r"r/forhire|reddit|freelancer(?:\.com)?)\s*$", re.I)


@dataclass
class Lead:
    job: Job
    kind: str = ""
    declared: str = ""
    summary: str = ""
    status: str = REVIEW
    note: str = ""
    verified: str = ""
    draft: Path | None = None


def clean_title(raw: str) -> str:
    title = htmllib.unescape(raw or "")
    title = re.sub(r"\s*-\s*Freelance Job\b.*$", "", title, flags=re.I)
    tail = re.search(r"\s+-\s+([A-Za-z ]{1,16})\s*(?:\.\.\.|…)?$", title)
    if tail and "freelance job in".startswith(tail.group(1).lower().strip()):
        title = title[:tail.start()]                       # Google cut "- Freelance Job in ..." short
    for _ in range(3):
        title = _SUFFIX.sub("", title).strip()
    title = re.sub(r"^\[hiring\]\s*", "", title, flags=re.I)
    return re.sub(r"\s*(?:\.\.\.|…)$", "", title).strip(" -|:")


def platform_job(item: dict, platform: str, page_text: str = "") -> Job:
    url = item.get("link") or item.get("url") or ""
    title = clean_title(item.get("title", ""))
    snippet = htmllib.unescape(str(item.get("snippet", "")))
    body = f"{item.get('title', '')}\n{snippet}\n{page_text}".strip()
    hr_email, alternates = pick_emails(body)
    return Job(source=platform, title=title, company=f"{platform} client", url=url, location="Remote",
               description=body[:15000], apply_url=url, hr_email=hr_email, alt_emails=alternates,
               date_posted=gs._age(str(item.get("date") or "")), date_trusted=True, is_remote=True)


# --------------------------------------------------------------------------- #
# Sources
# --------------------------------------------------------------------------- #
def freelancer_com(sess, cfg: dict) -> list[Lead]:
    """Freelancer.com public projects API: open projects tagged with DevOps / cloud skills."""
    src = cfg["sources"]["freelancer"]
    ids = src.get("skill_ids", [1678, 319, 1541, 1002, 2548, 1197, 2153, 2549, 1103, 1610, 1184])
    leads, now = [], time.time()
    for page in range(int(src.get("pages", 2))):
        try:
            r = sess.get("https://www.freelancer.com/api/projects/0.1/projects/active/",
                         params=[("jobs[]", i) for i in ids] + [("limit", 100), ("offset", page * 100),
                                                                ("full_description", "true"), ("job_details", "true"),
                                                                ("sort_field", "time_updated")])
            projects = r.json().get("result", {}).get("projects", [])
        except Exception as exc:  # noqa: BLE001
            log.info(f"  [freelancer.com] skipped: {type(exc).__name__}")
            break
        leads.extend(_freelancer_leads(projects, now))
        if len(projects) < 100:
            break
    seen_ids = {lead.job.external_id for lead in leads}
    for term in src.get("queries", []):
        try:
            r = sess.get("https://www.freelancer.com/api/projects/0.1/projects/active/",
                         params={"query": term, "limit": 50, "full_description": "true", "job_details": "true",
                                 "sort_field": "time_updated"})
            projects = r.json().get("result", {}).get("projects", [])
        except Exception as exc:  # noqa: BLE001
            log.info(f"  [freelancer.com] '{term}' skipped: {type(exc).__name__}")
            continue
        for p in projects:
            if str(p.get("id")) in seen_ids:
                continue
            seen_ids.add(str(p.get("id")))
            leads.extend(_freelancer_leads([p], now))
    log.info(f"  [freelancer.com] {len(leads)} open projects (DevOps / cloud skill tags + keyword searches)")
    return leads


def _freelancer_leads(projects: list[dict], now: float) -> list[Lead]:
    out = []
    for p in projects:
        if p.get("deleted") or p.get("nonpublic") or p.get("status") != "active" or p.get("language") not in (None, "en"):
            continue
        budget, cur = p.get("budget") or {}, (p.get("currency") or {}).get("code", "")
        skills = ", ".join(j.get("name", "") for j in p.get("jobs") or [])
        url = f"https://www.freelancer.com/projects/{p.get('seo_url')}"
        job = Job(source="Freelancer.com", title=p.get("title", ""), company="Freelancer.com client", url=url,
                  location="Remote", description=f"{p.get('description') or ''}\n\nSkills: {skills}",
                  apply_url=url, date_posted=to_datetime(p.get("time_submitted")), date_trusted=True,
                  is_remote=True, salary_min=budget.get("minimum"), salary_max=budget.get("maximum"),
                  salary_currency=cur, salary_interval="hourly" if p.get("type") == "hourly" else "fixed",
                  external_id=str(p.get("id")))
        age_h = (now - (p.get("time_submitted") or now)) / 3600
        bids = (p.get("bid_stats") or {}).get("bid_count", 0)
        out.append(Lead(job, declared=p.get("type", ""),
                        verified=f"Freelancer.com API: {p.get('frontend_project_status', 'open')} project, "
                                 f"{bids} bids, posted {age_h:.0f} h ago"))
    return out


_PPH_RELEVANT = re.compile(r"devops|aws|amazon|kubernetes|k8s|terraform|docker|cloud|ci-?cd|jenkins|linux|server|"
                           r"deploy|azure|gcp|helm|ansible|infrastructure|sre|monitoring|pipeline|nginx|lambda", re.I)


def peopleperhour(sess, cfg: dict) -> list[Lead]:
    """PeoplePerHour: search pages -> each project's own page, read from its JobPosting data."""
    src = cfg["sources"]["peopleperhour"]
    links: list[str] = []
    for kw in src.get("keywords", ["devops", "aws", "kubernetes", "terraform", "docker", "ci cd"]):
        try:
            page = sess.get("https://www.peopleperhour.com/freelance-jobs", params={"keyword": kw}).text
        except Exception as exc:  # noqa: BLE001
            log.info(f"  [peopleperhour] '{kw}' skipped: {type(exc).__name__}")
            continue
        for link in re.findall(r'href="(https://www\.peopleperhour\.com/freelance-jobs/[^"]+-\d+)"', page):
            if link not in links and _PPH_RELEVANT.search(link.rsplit("/", 1)[-1]):
                links.append(link)
    leads = []
    for link in links[: int(src.get("max_details", 40))]:
        try:
            page = sess.get(link).text
        except Exception:  # noqa: BLE001
            continue
        posting = None
        for block in re.findall(r'<script type="application/ld\+json">(.*?)</script>', page, re.S):
            try:
                data = json.loads(block)
            except ValueError:
                continue
            if isinstance(data, dict) and data.get("@type") == "JobPosting":
                posting = data
        if not posting:
            continue
        valid = to_datetime(posting.get("validThrough"))
        if valid and valid < datetime.now(timezone.utc):
            continue                                    # closed
        salary = posting.get("baseSalary") or {}
        value = salary.get("value") or {}
        amount = value.get("value") if isinstance(value, dict) else value
        try:
            amount = float(str(amount).replace(",", "")) if amount not in (None, "") else None
        except ValueError:
            amount = None
        unit = str(value.get("unitText", "") if isinstance(value, dict) else "").upper()
        job = Job(source="PeoplePerHour", title=htmllib.unescape(posting.get("title", "")),
                  company="PeoplePerHour client", url=link, location="Remote",
                  description=html_to_text(posting.get("description", "")), apply_url=link,
                  date_posted=to_datetime(posting.get("datePosted")), date_trusted=True, is_remote=True,
                  salary_min=amount, salary_max=amount, salary_currency=salary.get("currency", ""),
                  salary_interval="hourly" if unit == "HOUR" else "fixed")
        leads.append(Lead(job, declared="hourly" if unit == "HOUR" else "fixed",
                          verified=f"PeoplePerHour project page: posted {posting.get('datePosted', '?')[:10]}, "
                                   f"open until {str(posting.get('validThrough', '?'))[:10]}"))
    log.info(f"  [peopleperhour] {len(leads)} open projects (of {len(links)} relevant links)")
    return leads


def search_plan(cfg: dict, day: int) -> list[tuple[str, str]]:
    """[(kind, query)] for Google, past week. Upwork first: it has the most DevOps work."""
    g = cfg["sources"]["google"]
    terms = g.get("upwork_terms", [])
    start = day % max(1, len(terms))
    plan = [("platform", f"site:upwork.com/freelance-jobs/apply {terms[(start + i) % len(terms)]}")
            for i in range(min(int(g.get("upwork_searches", 8)), len(terms)))]
    skills = "(DevOps OR AWS OR Kubernetes OR Terraform OR Docker OR \"CI/CD\")"
    for site in g.get("platform_sites", []):
        plan.append(("platform", f"site:{site} {skills}"))
    places = g.get("post_regions", ["United States", "United Kingdom", "Europe", "Australia", "remote"])
    for i in range(int(g.get("linkedin_posts", 3))):
        plan.append(("post", f'("freelance" OR "part-time" OR "contract") ("DevOps" OR "AWS") ("looking for" OR "hiring") '
                             f'{gs.SHARE} (site:linkedin.com/posts OR site:linkedin.com/feed) {places[(day + i) % len(places)]}'))
    return plan[: int(g.get("max_searches", 20))]


def topup_plan(cfg: dict, used: set[str]) -> list[tuple[str, str]]:
    terms = [t for t in cfg["sources"]["google"].get("upwork_terms", []) + cfg["sources"]["google"].get("extra_terms", [])]
    return [("platform", q) for q in (f"site:upwork.com/freelance-jobs/apply {t}" for t in terms) if q not in used]


def run_searches(sess, plan: list[tuple[str, str]], cfg: dict, settings, seen_urls: set[str]) -> list[Lead]:
    providers = gs.available_providers()
    use_ddg = not providers and cfg["sources"]["duckduckgo"].get("enabled", True)
    if not providers and not use_ddg:
        log.info("  [google] no SERP key - marketplace searches skipped")
        return []
    leads = []
    for i, (kind, query) in enumerate(plan):
        try:
            items = gs.search_any(sess, query, providers, i, log.info) if providers else duckduckgo(sess, query)
            live = gs.available_providers()
            if not items and len(live) > 1:     # providers differ on site:/path queries - an empty answer gets a 2nd try
                items = gs.search_any(sess, query, live, i + 1, log.info)
        except Exception as exc:  # noqa: BLE001 - never print the exception text: URLs carry keys
            log.info(f"  [google] search failed: {type(exc).__name__} {getattr(getattr(exc, 'response', None), 'status_code', '')}")
            continue
        added = 0
        for item in items:
            url = item.get("link") or item.get("url") or ""
            key = canonical_url(url)
            if not url or key in seen_urls:
                continue
            seen_urls.add(key)
            host = urlsplit(url).netloc.lower().replace("www.", "").replace("old.", "")
            if kind == "post" or "linkedin.com" in host:
                job = gs._to_job("recruiter_posts", "Remote", "DevOps Engineer", item, gs._page_text(sess, url))
                if job:
                    job.source = "LinkedIn post"
                    leads.append(Lead(job, verified=f"LinkedIn post, Google index {item.get('date') or 'past week'}"))
                    added += 1
                continue
            platform = next(((name, rx) for dom, (name, rx) in PLATFORMS.items() if host.endswith(dom)), None)
            if platform is None or not platform[1].search(url):
                continue                                  # a listing page or an unknown site, not one project
            if platform[0] == "Reddit r/forhire" and not re.search(r"\[\s*hiring\s*\]", item.get("title", ""), re.I):
                continue                                  # "[For Hire]" = a freelancer advertising, not work offered
            page = "" if platform[0] in ("Upwork", "Guru", "Reddit r/forhire") else gs._page_text(sess, url)
            leads.append(Lead(platform_job(item, platform[0], page),
                              verified=f"{platform[0]} project page indexed by Google {item.get('date') or 'this week'}"
                                       + (", page read" if page else "")))
            added += 1
        log.info(f"  [google] {query[:70]}: {added} new")
    return leads


def duckduckgo(sess, query: str) -> list[dict]:
    """DuckDuckGo HTML results, past week - the fallback when no SERP key is configured."""
    from bs4 import BeautifulSoup

    r = sess.get("https://html.duckduckgo.com/html/", params={"q": query, "df": "w"})
    r.raise_for_status()
    soup = BeautifulSoup(r.text, "html.parser")
    out = []
    for res in soup.select(".result"):
        a, snip = res.select_one("a.result__a"), res.select_one(".result__snippet")
        if a and a.get("href", "").startswith("http"):
            out.append({"link": a["href"], "title": a.get_text(" ", strip=True),
                        "snippet": snip.get_text(" ", strip=True) if snip else ""})
    time.sleep(random.uniform(2, 4))
    return out


def from_boards(cfg: dict) -> list[Lead]:
    """LinkedIn / Indeed part-time and contract roles in the US, UK, EU, Australia, Canada and India."""
    from jobspy import scrape_jobs

    src = cfg["sources"]["boards"]
    per, limit = int(src.get("per_query", 20)), int(src.get("max_queries", 8))
    hours = int(cfg["max_age_days"]) * 24
    plan = []
    for kw in ("DevOps", "AWS DevOps", "Cloud Engineer"):
        plan += [("indeed", kw, None, True, "parttime", c) for c in ("usa", "uk", "australia", "canada", "germany", "india")]
        plan += [("linkedin", kw, loc, True, jt, "") for loc in ("United States", "United Kingdom", "European Union",
                                                               "Australia", "India") for jt in ("parttime", "contract")]
    random.Random(datetime.now(IST).toordinal()).shuffle(plan)
    leads = []
    for site, kw, location, remote, jtype, country in plan[:limit]:
        kwargs = dict(site_name=[site], search_term=kw, results_wanted=per, hours_old=hours, is_remote=remote,
                      job_type=jtype, description_format="markdown", verbose=0)
        if location:
            kwargs["location"] = location
        if site == "indeed":
            kwargs["country_indeed"] = country
        if site == "linkedin":
            kwargs["fetch_description"] = True     # full text: contact, skills, hours
        try:
            df = scrape_jobs(**kwargs)
        except Exception as exc:  # noqa: BLE001 - one site down never stops the run
            log.info(f"  [{site}] '{kw}' failed: {type(exc).__name__}")
            continue
        # is_remote False: remote must come from the posting itself, not from our search filter.
        q = boards.Query(site, kw, location, False, location or country.title(), 1)
        rows = df.to_dict("records") if len(df) else []
        for row in rows:
            job = boards._to_job(row, q)
            own_type = str(row.get("job_type") or "").lower()          # the posting's own employment type
            declared = "part_time" if "part" in own_type else "contract" if "contract" in own_type else ""
            if "full" in own_type and not declared:
                continue                                                 # a full-time job, whatever the filter said
            leads.append(Lead(job, declared=declared,
                              verified=f"{site.title()} listing ({own_type or 'type not stated'}), "
                                       f"posted {str(row.get('date_posted') or '?')[:10]}"))
        log.info(f"  [{site}] '{kw}' {location or country} {jtype}: {len(rows)} postings")
        time.sleep(random.uniform(2, 5))
    return leads


def from_feeds(sess) -> list[Lead]:
    from autopilot.discovery import remote_feeds as rf

    leads = []
    try:
        data = sess.get("https://remotive.com/api/remote-jobs", params={"category": "devops", "limit": 100}).json()
        for j in data.get("jobs", []):
            if j.get("job_type") in ("contract", "freelance", "part_time"):
                leads.append(Lead(rf._job("remotive", j.get("title"), j.get("company_name"), j.get("url"),
                                          j.get("candidate_required_location"), j.get("description"),
                                          j.get("publication_date"), external_id=j.get("id")),
                                  declared=j["job_type"], verified="Remotive API listing"))
    except Exception as exc:  # noqa: BLE001
        log.info(f"  [remotive] skipped: {type(exc).__name__}")
    try:
        data = sess.get("https://jobicy.com/api/v2/remote-jobs", params={"count": 50, "tag": "devops"}).json()
        for j in data.get("jobs", []):
            kinds = [str(t).lower() for t in (j.get("jobType") or [])]
            declared = next((k for k in kinds if k in ("contract", "part-time", "freelance")), "")
            if declared:
                leads.append(Lead(rf._job("jobicy", j.get("jobTitle"), j.get("companyName"), j.get("url"),
                                          j.get("jobGeo"), j.get("jobDescription") or j.get("jobExcerpt"),
                                          j.get("pubDate"), external_id=j.get("id")),
                                  declared=declared, verified="Jobicy API listing"))
    except Exception as exc:  # noqa: BLE001
        log.info(f"  [jobicy] skipped: {type(exc).__name__}")
    try:
        data = sess.get("https://himalayas.app/jobs/api", params={"limit": 100, "q": "devops"}).json()
        for j in data.get("jobs", []):
            kind = str(j.get("employmentType") or "").lower()
            declared = "contract" if "contract" in kind else "part_time" if "part" in kind else \
                "freelance" if "freelance" in kind else ""
            if declared:
                where = ", ".join(j.get("locationRestrictions") or []) or "Remote (anywhere)"
                leads.append(Lead(rf._job("himalayas", j.get("title"), j.get("companyName"),
                                          j.get("applicationLink") or j.get("guid"), where, j.get("description"),
                                          j.get("pubDate"), external_id=j.get("guid")),
                                  declared=declared, verified="Himalayas API listing"))
    except Exception as exc:  # noqa: BLE001
        log.info(f"  [himalayas] skipped: {type(exc).__name__}")
    log.info(f"  [feeds] {len(leads)} contract / freelance / part-time postings")
    return leads


# --------------------------------------------------------------------------- #
# Screening: genuine, current, side-work, 2-4 years, DevOps
# --------------------------------------------------------------------------- #
_DEVOPS = re.compile(r"devops|devsecops|\bsre\b|site reliability|\bcloud\b|\baws\b|amazon web services|azure|\bgcp\b|"
                     r"google cloud|kubernetes|\bk8s\b|docker|terraform|ansible|ci\s*/\s*cd|\bcicd\b|pipeline|jenkins|"
                     r"github actions|gitlab|helm|\beks\b|\becs\b|lambda|serverless|infrastructure|deploy|linux server|"
                     r"nginx|prometheus|grafana|monitoring|platform engineer|mlops|cloudformation|pulumi", re.I)
# A marketplace project's title must name DevOps / cloud work (their skill tags alone are not enough).
_TITLE_DEVOPS = re.compile(r"devops|devsecops|\bsre\b|site reliability|\bcloud\b|\baws\b|amazon web services|azure|"
                           r"\bgcp\b|google cloud|kubernetes|\bk8s\b|openshift|docker|container|terraform|ansible|"
                           r"cloudformation|pulumi|ci\s*/\s*cd|\bcicd\b|pipeline|jenkins|github actions|gitlab|argo\s*cd|"
                           r"helm|\beks\b|\becs\b|\baks\b|\bgke\b|lambda|serverless|infrastructure|deploy|\bvps\b|"
                           r"linux|nginx|server (?:setup|migration|configuration|admin)|prometheus|grafana|observability|"
                           r"\bmlops\b|platform engineer", re.I)
_SUPPORT_DESK = re.compile(r"\bl[123]\b|\bl[12]\s*/\s*l?[23]\b|help\s*desk|service desk|tech(?:nical)? support|"
                           r"desktop support|application support|information analyst|\[for hire\]", re.I)
_OFF_TOPIC = re.compile(r"wordpress|shopify|wix|webflow|landing page|logo|graphic|video|\bseo\b|marketing|"
                        r"data entry|virtual assistant|copywrit|translat|mobile app|android|\bios\b|flutter|"
                        r"react native|sales|accountant|tutor|photo|excel|crypto trading", re.I)
_FREELANCE_SCAM = re.compile(r"verified accounts?|accounts? (?:needed|wanted)|\bkyc\b|bypass|crack(?:ed)?\b|"
                             r"rent (?:your|an?) (?:account|laptop)|free (?:test|trial) task|unpaid|"
                             r"registration fee|pay (?:a|the) fee|telegram only|whatsapp only|airdrop|"
                             r"ghost ?writ|take (?:my|an?) (?:exam|test|interview)|residential prox|proxy rotation|"
                             r"proxy (?:an?|my|the) interview|interview proxy|fake (?:experience|resume)", re.I)


def screen(leads: list[Lead], cfg: dict, settings, profile, master: dict, history: dict,
           already: set[str] | None = None) -> tuple[list[Lead], Counter]:
    from autopilot.resume.parser import resume_text
    from autopilot.screening import eligibility_reason, experience_reason, genuineness, skill_match

    now = datetime.now(timezone.utc)
    resume_keys = implied(find_terms(resume_text(master)))
    cutoff = now - timedelta(days=int(cfg["max_age_days"]))
    seen_before = set(history.get("seen", {}))
    budget = cfg.get("budget", {})
    kept, why = [], Counter()
    urls, prints = set(already or ()), set()
    for lead in leads:
        job = lead.job
        key = canonical_url(job.url)
        title_key = re.sub(r"[^a-z0-9]+", " ", job.title.lower()).strip() + "|" + job.source
        text = f"{job.title}\n{job.description}"
        reason = ""
        if not job.url or key in urls or title_key in prints:
            reason = "duplicate"
        elif key in seen_before:
            reason = "already in an earlier run"
        elif job.date_posted and job.date_posted < cutoff:
            reason = "older than max_age_days"
        elif not job.date_posted and not job.date_trusted:
            reason = "no posting date"
        elif (job.source in PLATFORM_SOURCES or job.source == "LinkedIn post") and not _TITLE_DEVOPS.search(job.title):
            reason = "not DevOps / cloud work"      # project / post whose title is about something else
        elif _SUPPORT_DESK.search(job.title):
            reason = "not DevOps / cloud work"
        elif not _DEVOPS.search(job.title) and (_OFF_TOPIC.search(job.title) or len(set(find_terms(text)) & CORE_TERMS) < 2):
            reason = "not DevOps / cloud work"      # no DevOps word in the title: needs 2+ core DevOps skills in the text
        elif _SENIOR.search(job.title):
            reason = "senior / lead role (above 2-4 years)"
        elif _FREELANCE_SCAM.search(text):
            reason = "scam signal"
        if not reason:
            lead.kind = contract_type(job, lead.declared)
            if not lead.kind:
                reason = "not freelance / part-time / contract"
        if not reason:
            reason = side_time_reason(lead)
        if not reason:
            reason = genuineness(job, None, now) or experience_reason(job, settings)
        if not reason and job.source not in PLATFORM_SOURCES:
            reason = eligibility_reason(job, profile, settings)
            if not reason and not job.is_remote:
                reason = "on-site (side work must be remote)"
        if not reason and job.source in PLATFORM_SOURCES:
            top = usd_budget(job)
            hourly = job.salary_interval == "hourly" or "hourly" in lead.kind.lower()
            floor = float(budget.get("min_hourly_usd", 10) if hourly else budget.get("min_fixed_usd", 100))
            if top is not None and top < floor:
                reason = f"budget below ${floor:.0f}{'/hr' if hourly else ''}"
        if not reason:
            job.score, job.matched, job.missing = skill_match(job, resume_keys)
            if len(set(job.matched) | set(job.missing)) >= 3 and job.score < float(cfg["min_match"]):
                reason = "skill match too low"
            elif not job.matched:
                reason = "no matching skills"
        if reason:
            why[reason.split(" (")[0].split(":")[0]] += 1
            continue
        urls.add(key)
        prints.add(title_key)
        kept.append(lead)
    return kept, why


def ai_relevance(leads: list[Lead], cfg: dict, settings) -> tuple[list[Lead], int]:
    """Local Claude runs only (portal): Claude reads each kept post and drops the ones that are not
    really DevOps / cloud work for a 2-4 year engineer working 2-3 hours a day. GitHub runs skip this."""
    from autopilot.tailoring import claude_code
    from autopilot.tailoring.free_llm import QuickLLM

    if not claude_code.enabled() or not leads:
        return leads, 0
    llm = QuickLLM(settings, log.info, timeout=180, max_models=1, budget_seconds=600)
    availability = cfg.get("availability", "2-3 hours a day plus weekends")
    dropped: set[int] = set()
    for start in range(0, len(leads), 12):
        batch = leads[start:start + 12]
        listing = "\n\n".join(f"[{i}] {lead.job.title} ({lead.job.source})\n{lead.job.description[:1200]}"
                              for i, lead in enumerate(batch))
        data = llm.json(
            "You screen freelance leads for a DevOps / cloud engineer with 3 years of experience (AWS, Kubernetes, "
            f"Docker, Terraform, CI/CD, monitoring) who can work {availability}. For each post decide: is it mainly "
            "DevOps / cloud infrastructure / CI-CD / deployment / SRE work (not app, web, mobile or data-science "
            "development), right for 2-4 years of experience, and a genuine paid job (not a scam, not a freelancer "
            'advertising themselves)? Reply as JSON {"posts": [{"i": 0, "keep": true, "why": "few words"}]}.\n\n'
            + listing) or {}
        for item in data.get("posts") or []:
            try:
                idx = int(item.get("i"))
            except (TypeError, ValueError):
                continue
            if 0 <= idx < len(batch) and item.get("keep") is False:
                dropped.add(start + idx)
                batch[idx].note = f"dropped by Claude: {str(item.get('why') or '')[:80]}"
                log.info(f"  [claude] not a fit: {batch[idx].job.title[:60]} - {str(item.get('why') or '')[:60]}")
    return [lead for n, lead in enumerate(leads) if n not in dropped], len(dropped)


def order(leads: list[Lead]) -> list[Lead]:
    """Contactable first, then the freshest and best-matching."""
    now = datetime.now(timezone.utc)
    return sorted(leads, key=lambda lead: (not lead.job.hr_email,
                                           ((now - lead.job.date_posted).days if lead.job.date_posted else 7) // 2,
                                           -lead.job.score))


def summarize(leads: list[Lead], settings) -> None:
    """'Works / Description' in one line; Gemini/Groq/OpenRouter when available, else the post's first sentences.
    Only for your report - nothing an LLM writes here is sent to anyone."""
    from concurrent.futures import ThreadPoolExecutor

    from autopilot.tailoring.free_llm import QuickLLM

    # A convenience only: at most ~4 minutes, 40 s per call - unfinished ones keep the post's first sentences.
    llm = QuickLLM(settings, log.info, timeout=40, max_models=2, budget_seconds=240)
    for lead in leads:
        lead.summary = brief(lead.job.description) or lead.job.title
    if not llm:
        return

    def ask(batch):
        listing = "\n\n".join(f"[{i}] {lead.job.title}\n{lead.job.description[:1500]}" for i, lead in enumerate(batch))
        return batch, llm.json("For each work post below, write what has to be done in at most 25 plain words (the "
                               "project / tasks), using only what the post says. Reply as JSON "
                               '{"posts": [{"i": 0, "summary": "..."}]}\n\n' + listing) or {}

    started = time.time()
    batches = [leads[i:i + 10] for i in range(0, min(len(leads), 80), 10)]
    with ThreadPoolExecutor(max_workers=3) as pool:
        answers = list(pool.map(ask, batches))
    for batch, data in answers:
        for item in data.get("posts") or []:
            try:
                lead = batch[int(item.get("i"))]
            except (TypeError, ValueError, IndexError):
                continue
            text = re.sub(r"\s+", " ", str(item.get("summary") or "")).strip()
            if 5 <= len(text.split()) <= 40:
                lead.summary = text
    log.info(f"Descriptions: AI summaries for {sum(len(b) for b, d in answers if d)} leads ({time.time() - started:.0f}s)")


# --------------------------------------------------------------------------- #
# Emails
# --------------------------------------------------------------------------- #
def contacted_before(history: dict, days: int) -> set[str]:
    """Addresses this agent, the daily job autopilot or the older mailers already emailed."""
    out = set()
    cutoff = datetime.now(timezone.utc) - timedelta(days=days)
    for email, when in history.get("contacted", {}).items():
        if (to_datetime(when) or cutoff) >= cutoff:
            out.add(email.lower())
    db = APP / "data" / "autopilot.db"
    if db.exists():
        import sqlite3

        con = sqlite3.connect(str(db))
        try:
            out |= {r[0].lower() for r in con.execute("SELECT email FROM contacts WHERE last_contacted >= ?",
                                                      (cutoff.isoformat(),))}
        except sqlite3.Error:
            pass
        finally:
            con.close()
    from autopilot.apply import existing_pipeline_addresses

    return out | existing_pipeline_addresses(load_settings())


def draft(lead: Lead, cfg: dict, profile, master: dict, out_dir: Path) -> SimpleNamespace:
    from autopilot.resume.parser import resume_text
    from autopilot.tailoring import email_template, rules

    job = lead.job
    weights = rules.jd_weights(job)
    master_keys = implied(find_terms(resume_text(master)))
    known = rules._company_known(job) and not job.company.endswith(" client")
    kind = lead.kind.split(" - ")[0].lower()
    availability = cfg.get("availability", "2-3 hours a day plus weekends")
    intro = ("I'm reaching out about the {role} " + kind + " work{at}. I'm a {descriptor} with {years}+ years of "
             "hands-on experience in {skills}, and I can start {join}, working remotely " + availability + ".")
    text, html = email_template.build(
        job, profile, master, rules.matched_display(master_keys, weights, 10), rules.descriptor(master), known,
        weights, intro=intro,
        extra_facts=(("Availability", f"Freelance / part-time, remote - {availability}"),))
    subject = cfg["email"]["subject"].format(contract=lead.kind.split(" - ")[0], name=profile.name or "",
                                             title=email_template.display_title(job.title) or "DevOps Engineer",
                                             years=profile.years)
    folder = out_dir / "drafts" / f"{len(list((out_dir / 'drafts').glob('*'))) + 1:02d}-" \
                                  f"{re.sub(r'[^a-z0-9]+', '-', (job.company or 'lead').lower())[:30].strip('-')}"
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "email.txt").write_text(f"To: {job.hr_email}\nSubject: {subject}\n\n{text}", encoding="utf-8")
    (folder / "email.html").write_text(html, encoding="utf-8")
    lead.draft = folder
    resume = (HERE / cfg["email"]["resume"]).resolve()
    return SimpleNamespace(folder=folder, subject=subject, mail_text=text, mail_html=html, email_text=text,
                           email_html=html, resume_pdf=resume, attach_name=cfg["email"]["attachment_name"])


class Sender:
    def __init__(self):
        self.user, self.password = os.getenv("EMAIL_ADDRESS", ""), os.getenv("EMAIL_PASSWORD", "")
        self.host, self.port = os.getenv("SMTP_HOST", "smtp.gmail.com"), int(os.getenv("SMTP_PORT", "587"))
        self.smtp = None

    @property
    def ready(self) -> bool:
        return bool(self.user and self.password)

    def _connect(self):
        if self.smtp is None:
            self.smtp = smtplib.SMTP(self.host, self.port, timeout=30)
            self.smtp.ehlo()
            if os.getenv("SMTP_STARTTLS", "1") != "0":
                self.smtp.starttls(context=ssl.create_default_context())
                self.smtp.ehlo()
            self.smtp.login(self.user, self.password)
        return self.smtp

    def message(self, to: str, pkg, profile) -> EmailMessage:
        from email.utils import formataddr, make_msgid

        msg = EmailMessage()
        msg["From"] = formataddr((profile.name or "", self.user or profile.email or ""))
        msg["To"], msg["Subject"], msg["Date"] = to, pkg.subject, formatdate(localtime=True)
        msg["Message-ID"] = make_msgid(domain=(self.user or "localhost").split("@")[-1])
        if profile.email and profile.email.lower() != (self.user or "").lower():
            msg["Reply-To"] = profile.email
        msg.set_content(pkg.mail_text)
        msg.add_alternative(pkg.mail_html, subtype="html")
        if pkg.resume_pdf.exists():
            msg.add_attachment(pkg.resume_pdf.read_bytes(), maintype="application", subtype="pdf",
                               filename=pkg.attach_name)
        return msg

    def send(self, msg: EmailMessage) -> str:
        for attempt in (1, 2):
            try:
                self._connect().send_message(msg)
                return ""
            except smtplib.SMTPAuthenticationError as exc:
                raise RuntimeError(f"SMTP login failed ({exc.smtp_code}) - check EMAIL_PASSWORD") from exc
            except smtplib.SMTPRecipientsRefused:
                return "recipient refused"
            except (smtplib.SMTPServerDisconnected, ConnectionError, TimeoutError):
                self.smtp = None
                if attempt == 2:
                    return "SMTP connection lost"
            except smtplib.SMTPDataError as exc:
                if b"limit" in (exc.smtp_error or b"").lower():
                    raise RuntimeError("Gmail sending limit reached") from exc
                return f"SMTP error {exc.smtp_code}"
        return "not sent"

    def close(self):
        if self.smtp is not None:
            try:
                self.smtp.quit()
            except Exception:  # noqa: BLE001
                pass


# --------------------------------------------------------------------------- #
# Excel + report
# --------------------------------------------------------------------------- #
def contact_cell(job: Job) -> str:
    if job.hr_email:
        return job.hr_email
    return f"Apply on {job.source}" if job.source in PLATFORM_SOURCES else "Apply via the link"


def write_excel(leads: list[Lead], path: Path) -> Path:
    import pandas as pd
    from openpyxl.styles import Alignment, Font, PatternFill

    rows = [{"SR": i, "Company Name": lead.job.company or "Not stated", "Position": lead.job.title,
             "Works / Description": lead.summary, "Location": location_label(lead.job),
             "Salary / Rate": rate_of(lead.job), "Contract Type": lead.kind,
             "Contact / HR Email": contact_cell(lead.job),
             "Job Role": job_role(lead.job), "Status": lead.status}
            for i, lead in enumerate(leads, 1)]
    details = [{"SR": i, "Platform / source": lead.job.source, "Verified": lead.verified, "Job post": lead.job.url,
                "Apply / contact link": lead.job.apply_url or lead.job.url,
                "Posted": lead.job.date_posted.astimezone(IST).strftime("%d %b %Y") if lead.job.date_posted else "past week",
                "Skill match %": round(lead.job.score * 100), "Matched skills": ", ".join(display(lead.job.matched[:10])),
                "Other emails": ", ".join(lead.job.alt_emails), "Note": lead.note,
                "Draft": str(lead.draft.relative_to(path.parent)) if lead.draft else ""}
               for i, lead in enumerate(leads, 1)]
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame(rows, columns=COLUMNS).to_excel(xw, sheet_name="Leads", index=False)
        pd.DataFrame(details).to_excel(xw, sheet_name="Details", index=False)
        widths = {"SR": 5, "Company Name": 22, "Position": 36, "Works / Description": 60, "Location": 26,
                  "Salary / Rate": 22, "Contract Type": 20, "Contact / HR Email": 28, "Job Role": 34, "Status": 15}
        ws = xw.sheets["Leads"]
        for col, name in enumerate(COLUMNS, 1):
            cell = ws.cell(1, col)
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
            ws.column_dimensions[cell.column_letter].width = widths[name]
        fills = {SENT: "C6EFCE", DRAFTED: "FFEB9C", REVIEW: "DDEBF7"}
        for r, lead in enumerate(leads, 2):
            for col in range(1, len(COLUMNS) + 1):
                ws.cell(r, col).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(r, 3).hyperlink, ws.cell(r, 3).style = lead.job.url, "Hyperlink"
            ws.cell(r, 10).fill = PatternFill("solid", fgColor=fills.get(lead.status, "FFFFFF"))
        ws.freeze_panes, ws.auto_filter.ref = "A2", ws.dimensions
        d = xw.sheets["Details"]
        d.freeze_panes = "A2"
        for col, width in zip("ABCDEFGHIJK", (5, 18, 60, 50, 50, 12, 10, 40, 30, 40, 20)):
            d.column_dimensions[col].width = width
    return path


def email_report(xlsx: Path, leads: list[Lead], why: Counter, found: int, profile, health: str = "") -> None:
    from autopilot.report_mail import _table

    user, password = os.getenv("EMAIL_ADDRESS", ""), os.getenv("EMAIL_PASSWORD", "")
    to = os.getenv("REPORT_EMAIL") or profile.email or user
    if not (user and password and to):
        log.info("Report email skipped: EMAIL_ADDRESS / EMAIL_PASSWORD not set.")
        return
    c = Counter(lead.status for lead in leads)
    platforms = Counter(lead.job.source for lead in leads)
    msg = EmailMessage()
    msg["From"], msg["To"], msg["Date"] = user, to, formatdate(localtime=True)
    msg["Subject"] = (f"Freelance DevOps work {datetime.now(IST):%d %b %I:%M %p} IST - {len(leads)} verified leads, "
                      f"{c[SENT]} emailed")
    summary = (f"{found} posts checked, {len(leads)} verified freelance / part-time leads kept "
               f"({', '.join(f'{k} {v}' for k, v in platforms.most_common())}).\n"
               f"Emailed {c[SENT]}, drafted {c[DRAFTED]}, apply on the platform / review {c[REVIEW]}.\n"
               f"Dropped: {'; '.join(f'{k} {v}' for k, v in why.most_common(10))}\n{health}\n\nThe workbook is attached.")
    msg.set_content(summary)
    table = _table(["SR", "Platform", "Position", "Type", "Budget / rate", "Posted", "Contact", "Status", "Link"],
                   [[i, lead.job.source, lead.job.title, lead.kind, rate_of(lead.job),
                     lead.job.date_posted.astimezone(IST).strftime("%d %b") if lead.job.date_posted else "this week",
                     lead.job.hr_email or "apply on platform", lead.status, lead.job.url]
                    for i, lead in enumerate(leads, 1)], {8})
    msg.add_alternative(f'<div style="font-family:Arial,Helvetica,sans-serif;font-size:14px;color:#202124">'
                        f'<h2 style="margin:0 0 6px">Freelance DevOps / AWS work - side-work friendly, 2-4 yrs</h2>'
                        f'<p style="white-space:pre-line">{htmllib.escape(summary)}</p>{table}</div>', subtype="html")
    msg.add_attachment(xlsx.read_bytes(), maintype="application",
                       subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=xlsx.name)
    s = Sender()
    try:
        s.send(msg)
        log.info(f"Report emailed to {to}")
    finally:
        s.close()


# --------------------------------------------------------------------------- #
def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--send", action="store_true", help="email the contacts found (default: draft only)")
    ap.add_argument("--max-send", type=int, help="override email.max_send")
    ap.add_argument("--no-report", action="store_true", help="do not email the workbook to yourself")
    ap.add_argument("--config", default=str(HERE / "config.yaml"), help="settings file (default: config.yaml)")
    args = ap.parse_args(argv)

    cfg = yaml.safe_load(Path(args.config).read_text(encoding="utf-8"))
    stamp = datetime.now(IST).strftime("%Y%m%d_%H%M%S")
    run_dir = HERE / "runs" / stamp
    run_dir.mkdir(parents=True, exist_ok=True)
    setup_logging(run_dir)
    os.chdir(APP)
    load_env()
    settings, profile = load_settings(), load_profile()
    exp = cfg.get("experience", {})
    settings.data.setdefault("search", {}).update(
        max_age_days=int(cfg["max_age_days"]), experience_years=int(exp.get("years", 3)),
        experience_stretch=max(0, int(exp.get("max_years_asked", 4)) - int(exp.get("years", 3))))
    from autopilot.pipeline import get_master

    master = get_master(settings, log.info)
    hist_path = HERE / "data" / "history.json"
    history = json.loads(hist_path.read_text(encoding="utf-8")) if hist_path.exists() else {}
    log.info(f"Freelance work run {stamp} - {'SEND' if args.send else 'draft only'}")

    sess = ThrottledSession()
    src = cfg["sources"]
    leads: list[Lead] = []
    if src["freelancer"].get("enabled", True):
        leads += freelancer_com(sess, cfg)
    if src["peopleperhour"].get("enabled", True):
        leads += peopleperhour(sess, cfg)
    seen_urls: set[str] = set()
    plan = search_plan(cfg, datetime.now(IST).toordinal()) if src["google"].get("enabled", True) else []
    leads += run_searches(sess, plan, cfg, settings, seen_urls)
    if src["boards"].get("enabled", True):
        leads += from_boards(cfg)
    if src["feeds"].get("enabled", True):
        leads += from_feeds(sess)
    found = len(leads)
    all_found = list(leads)

    from autopilot.tailoring.free_llm import QuickLLM

    ai_search.read_posts(QuickLLM(settings, log.info), [lead.job for lead in leads], log.info)
    leads = [lead for lead in leads if not lead.job.title.endswith("[not a hiring post]")]
    kept, why = screen(leads, cfg, settings, profile, master, history)

    # Top up to min_leads with more Upwork skill searches while the search budget lasts.
    want, extra = int(cfg.get("min_leads", 15)), int(src["google"].get("topup_searches", 6))
    if len(kept) < want and extra and src["google"].get("enabled", True):
        more_plan = topup_plan(cfg, {q for _, q in plan})[:extra]
        log.info(f"Only {len(kept)} leads so far (target {want}) - {len(more_plan)} more searches")
        more = run_searches(sess, more_plan, cfg, settings, seen_urls)
        found += len(more)
        all_found += more
        extra_kept, extra_why = screen(more, cfg, settings, profile, master, history,
                                       already={canonical_url(lead.job.url) for lead in kept})
        kept += extra_kept
        why += extra_why
    kept, judged_out = ai_relevance(kept, cfg, settings)
    if judged_out:
        why["not a DevOps fit (Claude)"] += judged_out
    kept = order(kept)
    found_by = Counter(lead.job.source for lead in all_found)
    kept_by = Counter(lead.job.source for lead in kept)
    health = ("Per source (found -> kept): " + ", ".join(f"{k} {v} -> {kept_by.get(k, 0)}" for k, v in found_by.most_common())
              + f" | Search keys: {gs.provider_report()}")
    log.info(health)
    log.info(f"{found} posts -> {len(kept)} verified freelance / part-time leads "
             f"({', '.join(f'{k} {v}' for k, v in Counter(lead.job.source for lead in kept).most_common())}). "
             "Dropped: " + "; ".join(f"{k} {v}" for k, v in why.most_common(10)))
    summarize(kept, settings)

    done = contacted_before(history, int(cfg["email"]["recontact_after_days"]))
    sender = Sender()
    send = args.send and bool(cfg["email"].get("send", True))
    if send and not sender.ready:
        log.info("! EMAIL_ADDRESS / EMAIL_PASSWORD not set - drafting only")
        send = False
    cap, sent = args.max_send if args.max_send is not None else int(cfg["email"]["max_send"]), 0
    this_run: dict[str, int] = {}          # one email per address per run, even for two posts
    pause = cfg["email"].get("pause_seconds", [8, 20])
    try:
        for n, lead in enumerate(kept, 1):
            job = lead.job
            if not job.hr_email:
                lead.status = REVIEW
                lead.note = f"apply on {job.source} (link)" if job.source in PLATFORM_SOURCES else "apply via the link"
                continue
            if job.hr_email.lower() in done:
                lead.status, lead.note = REVIEW, "this address was already emailed (this agent or the job autopilot)"
                continue
            if job.hr_email.lower() in this_run:
                lead.status, lead.note = REVIEW, f"same contact as lead {this_run[job.hr_email.lower()]} this run"
                continue
            this_run[job.hr_email.lower()] = n
            if not mail_domain_ok(job.hr_email.split("@")[-1]):
                lead.status, lead.note = REVIEW, "email domain cannot receive mail"
                continue
            pkg = draft(lead, cfg, profile, master, run_dir)
            lead.status = DRAFTED
            if not send or sent >= cap:
                lead.note = "draft only" if not send else f"over this run's cap of {cap}"
                continue
            problem = sender.send(sender.message(job.hr_email, pkg, profile))
            if problem:
                lead.note = problem
                continue
            lead.status, sent = SENT, sent + 1
            done.add(job.hr_email.lower())
            history.setdefault("contacted", {})[job.hr_email.lower()] = datetime.now(timezone.utc).isoformat()
            log.info(f"  sent: {job.title} ({job.source})")
            time.sleep(random.uniform(float(pause[0]), float(pause[1])))
    except RuntimeError as exc:
        log.info(f"! sending stopped: {exc}")
    finally:
        sender.close()

    now = datetime.now(timezone.utc)
    seen = {k: v for k, v in history.get("seen", {}).items() if (to_datetime(v) or now) > now - timedelta(days=30)}
    if send:   # a draft-only run must not hide these leads from the next real run
        seen.update({canonical_url(lead.job.url): now.isoformat() for lead in kept})
    history["seen"] = seen
    hist_path.parent.mkdir(parents=True, exist_ok=True)
    hist_path.write_text(json.dumps(history, indent=1), encoding="utf-8")

    xlsx = write_excel(kept, run_dir / f"devops_leads_{stamp}.xlsx")
    c = Counter(lead.status for lead in kept)
    log.info(f"Done: {len(kept)} leads - Sent {c[SENT]}, Email Drafted {c[DRAFTED]}, Pending Review {c[REVIEW]}")
    log.info(f"Excel: {xlsx}")
    if os.getenv("GITHUB_STEP_SUMMARY"):
        with open(os.environ["GITHUB_STEP_SUMMARY"], "a", encoding="utf-8") as fh:
            fh.write(f"### Freelance DevOps work {stamp}\n- Posts checked: {found}\n- Verified leads: {len(kept)} "
                     f"({', '.join(f'{k} {v}' for k, v in Counter(lead.job.source for lead in kept).most_common())})\n"
                     f"- Emailed: {c[SENT]} · Drafted: {c[DRAFTED]} · Apply on platform / review: {c[REVIEW]}\n"
                     f"- {health}\n"
                     f"- Dropped: {'; '.join(f'{k} {v}' for k, v in why.most_common(10))}\n"
                     f"- Full workbook: in your email inbox\n")
    if cfg.get("report_email", True) and not args.no_report:
        try:
            email_report(xlsx, kept, why, found, profile, health)
        except Exception as exc:  # noqa: BLE001 - the run itself succeeded
            log.info(f"! report email failed: {type(exc).__name__}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
