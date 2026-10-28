"""python job_autopilot/.venv/bin/python -m pytest freelance_leads -q"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "job_autopilot"))

import freelance_leads as fl  # noqa: E402
from autopilot.models import Job  # noqa: E402


def job(**kw):
    base = dict(source="web:linkedin.com", title="DevOps Engineer", company="Acme", url="https://x.example/1",
                description="We need a DevOps engineer for AWS and Terraform work.")
    base.update(kw)
    return Job(**base)


def test_contract_type():
    assert fl.contract_type(job(title="DevOps Engineer (Contract, 6 months)")) == "Contract"
    assert fl.contract_type(job(description="Freelance DevOps help needed for a migration.")) == "Freelance"
    assert fl.contract_type(job(description="Part-time, 20 hours/week, remote.")) == "Part-Time"
    assert fl.contract_type(job(description="C2H role with a product company")) == "Contract"
    assert fl.contract_type(job(description="Full-time role. Your employment contract will be shared.")) == ""
    assert fl.contract_type(job(description="Write smart contract tooling")) == ""
    assert fl.contract_type(job(), declared="part_time") == "Part-Time"


def test_rate():
    assert fl.rate_of(job(description="Pay: ₹1,500 - 2,000 per hour, remote")) == "₹1,500 - 2,000 per hour"
    assert fl.rate_of(job(description="Budget 12-18 LPA, contract")) == "12-18 LPA"
    assert fl.rate_of(job(description="Rate $40/hr")) == "$40/hr"
    assert fl.rate_of(job()) == "Not stated"


def test_excel_columns(tmp_path):
    import openpyxl

    lead = fl.Lead(job(hr_email="hr@acme.com"), "Contract", summary="AWS migration", status=fl.SENT)
    path = fl.write_excel([lead], tmp_path / "devops_leads_x.xlsx")
    ws = openpyxl.load_workbook(path)["Leads"]
    assert [c.value for c in ws[1]] == fl.COLUMNS
    assert [c.value for c in ws[2]][:3] == [1, "Acme", "DevOps Engineer"] and ws.cell(2, 10).value == "Sent"


# --------------------------------------------------------------------------- freelance marketplaces
import json  # noqa: E402
import time  # noqa: E402
from types import SimpleNamespace  # noqa: E402


def test_clean_title_and_platform_job():
    assert fl.clean_title("AWS DevOps for Logistics Backend - Freelance Job in DevOps Engineering - $25.00-$50.00 "
                          "Hourly - Upwork") == "AWS DevOps for Logistics Backend"
    assert fl.clean_title("[Hiring] Terraform help needed : r/forhire") == "Terraform help needed"
    j = fl.platform_job({"link": "https://www.upwork.com/freelance-jobs/apply/DevOps_~0221/",
                         "title": "Kubernetes cluster setup - Upwork", "date": "3 days ago",
                         "snippet": "Hourly: $25.00-$40.00. Set up EKS with Terraform and Helm."}, "Upwork")
    assert j.title == "Kubernetes cluster setup" and j.company == "Upwork client" and j.is_remote
    assert j.date_posted is not None and fl.rate_of(j) == "$25.00-$40.00"


def test_side_time_and_seniority():
    project = fl.Lead(job(source="Upwork", description="Fixed-price: migrate our app to AWS ECS, Terraform."))
    assert fl.side_time_reason(project) == ""
    full = fl.Lead(job(source="Upwork", description="Full-time, 40 hours/week, Monday to Friday DevOps role."))
    assert fl.side_time_reason(full) == "full-time commitment"
    contract = fl.Lead(job(source="linkedin", description="12 month contract, full-time, AWS DevOps."), kind="Contract")
    assert fl.side_time_reason(contract).startswith("full-time contract")
    part = fl.Lead(job(source="linkedin", description="Contract, 10-15 hours/week, AWS + Terraform."), kind="Contract")
    assert fl.side_time_reason(part) == ""
    assert fl._SENIOR.search("Senior DevOps Engineer") and not fl._SENIOR.search("DevOps Engineer (2-4 yrs)")


class FakeSession:
    def __init__(self, routes):
        self.routes = routes

    def get(self, url, params=None, **kw):
        for prefix, body in self.routes.items():
            if url.startswith(prefix):
                text = body if isinstance(body, str) else json.dumps(body)
                return SimpleNamespace(text=text, json=lambda t=text: json.loads(t), ok=True, status_code=200)
        return SimpleNamespace(text="", json=lambda: {}, ok=False, status_code=404)


def test_freelancer_api_parsing():
    now = time.time()
    api = {"result": {"projects": [
        {"id": 1, "title": "Terraform modules for AWS", "status": "active", "language": "en", "type": "fixed",
         "seo_url": "terraform/Terraform-modules", "currency": {"code": "USD"}, "budget": {"minimum": 250, "maximum": 750},
         "description": "Write Terraform for VPC, EKS.", "jobs": [{"name": "Terraform"}, {"name": "Amazon Web Services"}],
         "time_submitted": now - 3600, "bid_stats": {"bid_count": 12}, "frontend_project_status": "open"},
        {"id": 2, "title": "Closed one", "status": "closed", "language": "en", "type": "fixed", "seo_url": "x",
         "currency": {"code": "USD"}, "budget": {}, "time_submitted": now}]}}
    leads = fl.freelancer_com(FakeSession({"https://www.freelancer.com/api": api}),
                              {"sources": {"freelancer": {"pages": 1}}})
    assert len(leads) == 1
    lead = leads[0]
    assert lead.job.url == "https://www.freelancer.com/projects/terraform/Terraform-modules"
    assert fl.rate_of(lead.job) == "USD 250-750 fixed price" and "12 bids" in lead.verified


def test_peopleperhour_reads_the_project_page():
    listing = '<a href="https://www.peopleperhour.com/freelance-jobs/technology-programming/devops/aws-devops-setup-4528561">x</a>' \
              '<a href="https://www.peopleperhour.com/freelance-jobs/design/logo-design-4528562">y</a>'
    posting = {"@type": "JobPosting", "title": "AWS DevOps setup", "description": "<p>Set up CI/CD and ECS</p>",
               "datePosted": "2099-01-01 10:00:00", "validThrough": "2099-02-01 10:00:00",
               "baseSalary": {"currency": "USD", "value": {"unitText": "HOUR", "value": "30"}}}
    page = f'<script type="application/ld+json">{json.dumps(posting)}</script>'
    sess = FakeSession({"https://www.peopleperhour.com/freelance-jobs/technology": page,
                        "https://www.peopleperhour.com/freelance-jobs": listing})
    leads = fl.peopleperhour(sess, {"sources": {"peopleperhour": {"keywords": ["devops"], "max_details": 5}}})
    assert len(leads) == 1 and leads[0].job.title == "AWS DevOps setup"          # the logo job was not fetched
    assert fl.rate_of(leads[0].job) == "USD 30 /hr" and "open until 2099-02-01" in leads[0].verified


def test_marketplace_titles_must_name_devops_work():
    assert fl.clean_title("Cloud Security Integration Specialist - Fr") == "Cloud Security Integration Specialist"
    assert fl.clean_title("Platform DevOps Engineer - Freelance Job") == "Platform DevOps Engineer"
    for title in ("Complete PHP Payment Gateway Platform", "React & Python Business Dashboard",
                  "Secure Availability Monitoring And Booking", "Tech Talent Sourcing Partner"):
        assert not fl._TITLE_DEVOPS.search(title), title
    for title in ("AWS AI API Infrastructure Optimization", "ProCall Hub Deployment", "Kubernetes cluster setup",
                  "Terraform modules for AWS", "CI/CD pipeline with GitHub Actions"):
        assert fl._TITLE_DEVOPS.search(title), title
    assert fl._FREELANCE_SCAM.search("extract the feed using residential proxies")


def test_support_desks_and_for_hire_posts_are_not_work():
    for title in ("Information Analyst – Application L1 Support", "Remote Tech Support (L1-L3), Firewalls, DevOps",
                  "[For Hire] Devops engineer and cloud costs optimization"):
        assert fl._SUPPORT_DESK.search(title), title
    assert not fl._SUPPORT_DESK.search("DevOps Engineer - Kubernetes support for our platform team")
