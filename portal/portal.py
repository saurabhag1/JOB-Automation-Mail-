#!/usr/bin/env python3
"""Local control center for the job-automation workflows.

    python3 portal/portal.py            open http://127.0.0.1:8765 in your browser
    python3 portal/portal.py --port 9000 --no-browser

Shows every GitHub workflow (schedule in IST, what it does, its latest GitHub runs) and runs the
two that are safe to run locally - Daily Job Autopilot and Freelance / Contract DevOps Leads -
on this Mac with Claude (through the Claude Code CLI) instead of the free models. Local runs:
  * use Claude Code for tailoring and the AI search helpers (AUTOPILOT_LLM=claude-code);
  * never re-email anyone your Gmail account emailed in the last 60 days - the GitHub runs
    included - by reading the recipients in Gmail's Sent folder (AUTOPILOT_GMAIL_DEDUPE=1);
  * email you the report at the end, like the GitHub runs.
The GitHub scheduled runs are not touched: they never set those variables.

The server listens on 127.0.0.1 only and every API call needs the token embedded in the page.
"""

from __future__ import annotations

import os
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
APP = REPO / "job_autopilot"
VENV = APP / ".venv"
if os.path.realpath(sys.prefix) != os.path.realpath(VENV) and (VENV / "bin" / "python").exists():
    os.execv(str(VENV / "bin" / "python"), [str(VENV / "bin" / "python"), str(Path(__file__).resolve()), *sys.argv[1:]])
sys.path.insert(0, str(APP))

import argparse  # noqa: E402
import csv  # noqa: E402
import json  # noqa: E402
import re  # noqa: E402
import secrets  # noqa: E402
import shutil  # noqa: E402
import signal  # noqa: E402
import subprocess  # noqa: E402
import threading  # noqa: E402
import time  # noqa: E402
import webbrowser  # noqa: E402
from datetime import datetime, timedelta, timezone  # noqa: E402
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer  # noqa: E402
from urllib.parse import parse_qs, urlsplit  # noqa: E402

import requests  # noqa: E402
import yaml  # noqa: E402

IST = timezone(timedelta(hours=5, minutes=30), "IST")
DATA = HERE / "data"
RUNS = DATA / "runs"
SETTINGS_FILE = DATA / "settings.json"
PY = str(VENV / "bin" / "python")
TOKEN = secrets.token_urlsafe(24)
DEFAULTS = {"use_claude": True, "model": "opus", "effort": "high", "gmail_dedupe": True,
            "autopilot_max": 40, "freelance_max_send": 30}
SERP_KEYS = ["SERPER_KEY", "SEARCHAPI_KEY", "SCRAPERAPI_KEY", "SERPAPI_KEY", "SCRAPINGDOG_KEY"]
LLM_KEYS = ["GEMINI_API_KEY", "GROQ_API_KEY", "OPENROUTER_API_KEY"]

# Workflows that are safe to run on this Mac. The bulk mailer, its follow-ups and the lead
# scraper commit their state back to GitHub; running them here too would send duplicates.
LOCAL = {
    "daily-job-autopilot.yml": {
        "key": "autopilot",
        "local_note": "Finds jobs, tailors your resume with Claude, emails HR, submits Greenhouse / Lever / Ashby "
                      "forms, then emails you the report.",
        "steps": lambda o: [[PY, "job_hunt.py", "--unattended", "--headless", "--max", str(o["autopilot_max"])]
                            + (["--dry-run"] if o["dry_run"] else []),
                            [PY, "job_hunt.py", "mail-report"]],
    },
    "freelance-leads.yml": {
        "key": "freelance",
        "local_note": "Finds freelance / part-time DevOps work worldwide (Upwork, Freelancer.com, PeoplePerHour, ...), "
                      "emails contacts, and emails you the Excel.",
        "steps": lambda o: [[PY, "freelance_leads/freelance_leads.py", "--max-send", str(o["freelance_max_send"])]
                            + ([] if o["dry_run"] else ["--send"])],
    },
}
GITHUB_ONLY_NOTE = ("Runs on GitHub only: it commits its own state back to the repository, so a second, local "
                    "copy would email the same people twice. Use 'Open on GitHub' to run it there.")


# --------------------------------------------------------------------------- #
# Settings, environment, status
# --------------------------------------------------------------------------- #
def load_settings() -> dict:
    try:
        return {**DEFAULTS, **json.loads(SETTINGS_FILE.read_text(encoding="utf-8"))}
    except (OSError, ValueError):
        return dict(DEFAULTS)


def save_settings(data: dict) -> dict:
    current = load_settings()
    for key, value in data.items():
        if key in DEFAULTS:
            current[key] = type(DEFAULTS[key])(value) if not isinstance(DEFAULTS[key], bool) else bool(value)
    if current["effort"] not in ("low", "medium", "high", "xhigh", "max"):
        current["effort"] = "high"
    current["model"] = re.sub(r"[^A-Za-z0-9._\[\]-]", "", current["model"]) or "opus"
    DATA.mkdir(parents=True, exist_ok=True)
    SETTINGS_FILE.write_text(json.dumps(current, indent=1), encoding="utf-8")
    return current


def keys_present() -> dict:
    """Which keys are configured (names only - never values)."""
    from autopilot.config import load_env

    load_env()
    return {"gmail": bool(os.getenv("EMAIL_ADDRESS") and os.getenv("EMAIL_PASSWORD")),
            "search": [k.replace("_KEY", "").lower() for k in SERP_KEYS if os.getenv(k)],
            "free_llms": [k.split("_")[0].lower() for k in LLM_KEYS if os.getenv(k)],
            "github_token": bool(os.getenv("GITHUB_TOKEN") or os.getenv("GH_TOKEN"))}


def repo_slug() -> str:
    try:
        url = subprocess.run(["git", "-C", str(REPO), "remote", "get-url", "origin"], capture_output=True,
                             text=True, timeout=10).stdout.strip()
    except (OSError, subprocess.SubprocessError):
        return ""
    m = re.search(r"github\.com[:/](.+?)(?:\.git)?$", url)
    return m.group(1) if m else ""


SLUG = repo_slug()


def cron_ist(expr: str) -> str:
    """'30 5 * * *' (UTC) -> 'Every day, 11:00 AM IST'."""
    parts = expr.split()
    if len(parts) != 5 or not parts[0].isdigit() or not parts[1].isdigit():
        return f"cron {expr} (UTC)"
    utc = datetime(2026, 1, 5, int(parts[1]), int(parts[0]), tzinfo=timezone.utc)   # a Monday
    ist = utc.astimezone(IST)
    days = {"*": "Every day", "1-5": "Mon-Fri", "1-6": "Mon-Sat", "0,6": "Weekends"}.get(parts[4], f"days {parts[4]}")
    if parts[4] == "1-5" and ist.day != utc.day:
        days = "Mon-Fri (IST date)"
    return f"{days}, {ist.strftime('%I:%M %p').lstrip('0')} IST"


def workflows() -> list[dict]:
    out = []
    for path in sorted((REPO / ".github" / "workflows").glob("*.yml")):
        text = path.read_text(encoding="utf-8")
        try:
            doc = yaml.safe_load(text) or {}
        except yaml.YAMLError:
            continue
        trigger = doc.get("on") or doc.get(True) or {}
        crons = [c.get("cron") for c in (trigger.get("schedule") or []) if isinstance(c, dict)] if isinstance(trigger, dict) else []
        after = (trigger.get("workflow_run") or {}).get("workflows") if isinstance(trigger, dict) else None
        lines, started = [], False
        for line in text.splitlines():
            if line.startswith("#"):
                started = True
                lines.append(line.lstrip("# ").rstrip())
            elif started and line.strip() and not line.startswith("#"):
                if line.startswith("on:"):
                    break
        about = " ".join(x for x in lines if x and not re.match(r"^[A-Z_]+(,|\s{2})", x))[:700]
        schedule = ", ".join(cron_ist(c) for c in crons if c) or (f"After '{after[0]}' finishes" if after else "Manual")
        local = LOCAL.get(path.name)
        out.append({"file": path.name, "name": doc.get("name", path.stem), "schedule": schedule,
                    "about": about or "", "local": bool(local), "key": local["key"] if local else "",
                    "local_note": local["local_note"] if local else GITHUB_ONLY_NOTE,
                    "github_url": f"https://github.com/{SLUG}/actions/workflows/{path.name}" if SLUG else ""})
    return out


_gh_cache: dict[str, tuple[float, list]] = {}


def github_runs(file: str) -> list[dict]:
    """Latest runs of a workflow from GitHub's public API (cached for 2 minutes)."""
    if not SLUG:
        return []
    hit = _gh_cache.get(file)
    if hit and time.time() - hit[0] < 120:
        return hit[1]
    runs = []
    try:
        headers = {"Accept": "application/vnd.github+json"}
        token = os.getenv("GITHUB_TOKEN") or os.getenv("GH_TOKEN")
        if token:
            headers["Authorization"] = f"Bearer {token}"
        r = requests.get(f"https://api.github.com/repos/{SLUG}/actions/workflows/{file}/runs",
                         params={"per_page": 6}, headers=headers, timeout=15)
        for run in r.json().get("workflow_runs", []) if r.ok else []:
            created = datetime.fromisoformat(run["created_at"].replace("Z", "+00:00")).astimezone(IST)
            runs.append({"status": run.get("status"), "conclusion": run.get("conclusion"), "event": run.get("event"),
                         "when": created.strftime("%d %b, %I:%M %p"), "url": run.get("html_url")})
    except (requests.RequestException, ValueError, KeyError):
        pass
    _gh_cache[file] = (time.time(), runs)
    return runs


def dispatch_github(file: str) -> str:
    token = os.getenv("GITHUB_TOKEN") or os.getenv("GH_TOKEN")
    if not token or not SLUG:
        return "Add GITHUB_TOKEN (a token with 'actions: write') to job_autopilot/.env to start GitHub runs from here."
    r = requests.post(f"https://api.github.com/repos/{SLUG}/actions/workflows/{file}/dispatches",
                      headers={"Authorization": f"Bearer {token}", "Accept": "application/vnd.github+json"},
                      json={"ref": "main"}, timeout=20)
    _gh_cache.pop(file, None)
    return "Started on GitHub." if r.status_code == 204 else f"GitHub refused ({r.status_code}): {r.text[:160]}"


# --------------------------------------------------------------------------- #
# Local runs
# --------------------------------------------------------------------------- #
class Runner:
    def __init__(self):
        self.lock = threading.Lock()
        self.active: dict | None = None
        self.proc: subprocess.Popen | None = None
        RUNS.mkdir(parents=True, exist_ok=True)

    def history(self) -> list[dict]:
        out = []
        for meta in sorted(RUNS.glob("*/meta.json"), reverse=True)[:30]:
            try:
                out.append(json.loads(meta.read_text(encoding="utf-8")))
            except (OSError, ValueError):
                continue
        return out

    def start(self, file: str, dry_run: bool) -> dict:
        job = LOCAL.get(file)
        if not job:
            raise ValueError("this workflow runs on GitHub only")
        with self.lock:
            if self.active and self.active["status"] == "running":
                raise ValueError(f"'{self.active['title']}' is still running - wait or stop it first")
            opts = {**load_settings(), "dry_run": dry_run}
            run_id = datetime.now(IST).strftime("%Y%m%d-%H%M%S") + f"-{job['key']}"
            folder = RUNS / run_id
            folder.mkdir(parents=True)
            title = next((w["name"] for w in workflows() if w["file"] == file), file)
            self.active = {"id": run_id, "file": file, "key": job["key"], "title": title, "status": "running",
                           "mode": "dry run" if dry_run else "live (sends email)",
                           "llm": f"Claude Code {opts['model']} ({opts['effort']})" if opts["use_claude"] else "free LLMs",
                           "started": datetime.now(IST).strftime("%d %b %Y, %I:%M:%S %p"), "ended": "", "exit": None}
            self._save()
            env = {**os.environ, "PYTHONUNBUFFERED": "1"}
            if opts["use_claude"]:
                env.update(AUTOPILOT_LLM="claude-code", AUTOPILOT_CLAUDE_MODEL=opts["model"],
                           AUTOPILOT_CLAUDE_EFFORT=opts["effort"])
            if opts["gmail_dedupe"]:
                env["AUTOPILOT_GMAIL_DEDUPE"] = "1"
            env.pop("CI", None)
            threading.Thread(target=self._run, args=(job["steps"](opts), env, folder), daemon=True).start()
            return self.active

    def _save(self):
        (RUNS / self.active["id"] / "meta.json").write_text(json.dumps(self.active, indent=1), encoding="utf-8")

    def _run(self, steps: list[list[str]], env: dict, folder: Path):
        log = open(folder / "log.txt", "a", encoding="utf-8", buffering=1)
        code = 0
        try:
            for cmd in steps:
                if self.active.get("status") == "stopped":
                    break
                log.write(f"\n$ {' '.join(Path(c).name if i == 0 else c for i, c in enumerate(cmd))}\n")
                if shutil.which("caffeinate"):      # keep the Mac awake until this step ends
                    cmd = ["caffeinate", "-i", *cmd]
                self.proc = subprocess.Popen(cmd, cwd=str(REPO), env=env, stdout=log, stderr=subprocess.STDOUT,
                                             start_new_session=True)
                code = self.proc.wait()
                if code != 0 and self.active.get("status") != "stopped":
                    log.write(f"\n[step exited with code {code}]\n")
        except Exception as exc:  # noqa: BLE001
            log.write(f"\n[portal error: {type(exc).__name__}: {exc}]\n")
            code = 1
        finally:
            log.close()
            with self.lock:
                if self.active.get("status") != "stopped":
                    self.active["status"] = "finished" if code == 0 else "failed"
                self.active["ended"] = datetime.now(IST).strftime("%d %b %Y, %I:%M:%S %p")
                self.active["exit"] = code
                self._save()
                self.proc = None

    def stop(self) -> str:
        with self.lock:
            if not self.active or self.active["status"] != "running":
                return "Nothing is running."
            self.active["status"] = "stopped"
            self._save()
            if self.proc and self.proc.poll() is None:
                try:
                    os.killpg(self.proc.pid, signal.SIGTERM)
                except (ProcessLookupError, PermissionError):
                    pass
        return "Stopping."

    def log(self, run_id: str, offset: int) -> tuple[str, int]:
        path = RUNS / re.sub(r"[^0-9A-Za-z_-]", "", run_id) / "log.txt"
        if not path.exists():
            return "", offset
        with open(path, "rb") as fh:
            fh.seek(offset)
            chunk = fh.read(200_000)
        return chunk.decode("utf-8", "replace"), offset + len(chunk)


RUNNER = Runner()


# --------------------------------------------------------------------------- #
# Results
# --------------------------------------------------------------------------- #
_ITEM = re.compile(r"^- \*\*(.+?):\*\* (.*)$", re.M)


def _rows(path: Path) -> list[dict]:
    try:
        with open(path, newline="", encoding="utf-8") as fh:
            return list(csv.DictReader(fh))
    except OSError:
        return []


def autopilot_results() -> dict:
    runs = sorted((p for p in (REPO / "job_runs").glob("20*_IST*") if p.is_dir()), key=lambda p: p.stat().st_mtime,
                  reverse=True)
    if not runs:
        return {}
    run = runs[0]
    md = sorted(run.glob("*.md"))
    items = dict(_ITEM.findall(md[0].read_text(encoding="utf-8"))) if md else {}
    sent = [r for r in _rows(run / "sent.csv") if r.get("Channel") != "manual"]
    manual = _rows(run / "apply_manually.csv")
    xlsx = sorted(run.glob("*.xlsx"))
    return {"folder": run.name, "items": items,
            "sent": [{k: r.get(k, "") for k in ("Time (IST)", "Channel", "Sent to / where", "Company", "Job title",
                                                "Status", "Job post")} for r in sent][:80],
            "manual": [{k: r.get(k, "") for k in ("Job title", "Company", "Location", "Why by hand", "Apply here")}
                       for r in manual][:80],
            "excel": str(xlsx[0].relative_to(REPO)) if xlsx else "",
            "history": [p.name for p in runs[:10]]}


def freelance_results() -> dict:
    files = sorted((REPO / "freelance_leads" / "runs").glob("*/devops_leads_*.xlsx"), key=lambda p: p.stat().st_mtime,
                   reverse=True)
    if not files:
        return {}
    from openpyxl import load_workbook

    wb = load_workbook(files[0], read_only=True)
    leads = list(wb["Leads"].iter_rows(values_only=True))
    details = {r[0]: r for r in wb["Details"].iter_rows(min_row=2, values_only=True)} if "Details" in wb.sheetnames else {}
    head = [str(h) for h in leads[0]] if leads else []
    dhead = [str(c.value) for c in next(wb["Details"].iter_rows(max_row=1))] if "Details" in wb.sheetnames else []
    rows = []
    for row in leads[1:]:
        item = dict(zip(head, row))
        extra = dict(zip(dhead, details.get(row[0], ())))
        item.update({"Platform": extra.get("Platform / source", ""), "Verified": extra.get("Verified", ""),
                     "Posted": extra.get("Posted", ""), "Link": extra.get("Job post", "")})
        rows.append({k: ("" if v is None else str(v)) for k, v in item.items()})
    wb.close()
    return {"file": str(files[0].relative_to(REPO)), "run": files[0].parent.name, "leads": rows}


# --------------------------------------------------------------------------- #
# HTTP
# --------------------------------------------------------------------------- #
class Handler(BaseHTTPRequestHandler):
    server_version = "JobPortal/1.0"

    def log_message(self, fmt, *args):   # quiet console
        pass

    def _send(self, code: int, body: bytes, ctype: str, extra: dict | None = None):
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        for k, v in (extra or {}).items():
            self.send_header(k, v)
        self.end_headers()
        self.wfile.write(body)

    def _json(self, data, code: int = 200):
        self._send(code, json.dumps(data).encode(), "application/json")

    def _allowed(self) -> bool:
        host = (self.headers.get("Host") or "").split(":")[0]
        if host not in ("127.0.0.1", "localhost"):
            return False                                        # DNS-rebinding guard
        url = urlsplit(self.path)
        if url.path == "/" or url.path.startswith("/static/"):
            return True
        token = self.headers.get("X-Portal-Token") or parse_qs(url.query).get("token", [""])[0]
        return secrets.compare_digest(token, TOKEN)

    def do_GET(self):
        if not self._allowed():
            return self._send(403, b"forbidden", "text/plain")
        url = urlsplit(self.path)
        q = parse_qs(url.query)
        if url.path == "/":
            html = (HERE / "static" / "index.html").read_text(encoding="utf-8").replace("__PORTAL_TOKEN__", TOKEN)
            return self._send(200, html.encode(), "text/html; charset=utf-8")
        if url.path == "/api/state":
            flows = workflows()
            for w in flows:
                w["github_runs"] = github_runs(w["file"])
            return self._json({"workflows": flows, "settings": load_settings(), "keys": keys_present(),
                               "claude_cli": bool(shutil.which("claude")), "active": RUNNER.active,
                               "history": RUNNER.history(), "repo": SLUG})
        if url.path == "/api/log":
            text, offset = RUNNER.log(q.get("id", [""])[0], int(q.get("offset", ["0"])[0] or 0))
            return self._json({"text": text, "offset": offset})
        if url.path == "/api/results":
            return self._json({"autopilot": autopilot_results(), "freelance": freelance_results()})
        if url.path == "/files":
            rel = q.get("path", [""])[0]
            path = (REPO / rel).resolve()
            roots = [(REPO / "job_runs").resolve(), (REPO / "freelance_leads" / "runs").resolve()]
            if not any(str(path).startswith(str(r) + os.sep) for r in roots) or not path.is_file():
                return self._send(404, b"not found", "text/plain")
            ctype = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" if path.suffix == ".xlsx" \
                else "text/plain; charset=utf-8"
            return self._send(200, path.read_bytes(), ctype, {"Content-Disposition": f'attachment; filename="{path.name}"'})
        return self._send(404, b"not found", "text/plain")

    def do_POST(self):
        if not self._allowed():
            return self._send(403, b"forbidden", "text/plain")
        length = int(self.headers.get("Content-Length") or 0)
        try:
            body = json.loads(self.rfile.read(length) or b"{}")
        except ValueError:
            body = {}
        path = urlsplit(self.path).path
        try:
            if path == "/api/run":
                return self._json({"ok": True, "active": RUNNER.start(body.get("file", ""), bool(body.get("dry_run")))})
            if path == "/api/stop":
                return self._json({"ok": True, "message": RUNNER.stop()})
            if path == "/api/settings":
                return self._json({"ok": True, "settings": save_settings(body)})
            if path == "/api/github-dispatch":
                return self._json({"ok": True, "message": dispatch_github(body.get("file", ""))})
            if path == "/api/claude-test":
                return self._json(claude_test())
        except ValueError as exc:
            return self._json({"ok": False, "message": str(exc)}, 409)
        return self._send(404, b"not found", "text/plain")


def claude_test() -> dict:
    s = load_settings()
    os.environ.update(AUTOPILOT_LLM="claude-code", AUTOPILOT_CLAUDE_MODEL=s["model"], AUTOPILOT_CLAUDE_EFFORT=s["effort"])
    from autopilot.tailoring import claude_code

    started = time.time()
    try:
        data = claude_code.ask('Reply with JSON {"ok": true, "model": "<your model name>"}.', timeout=120)
        return {"ok": True, "message": f"Claude answered in {time.time() - started:.0f}s: {json.dumps(data)}"}
    except claude_code.ClaudeCodeError as exc:
        return {"ok": False, "message": f"Claude Code failed: {exc}"}
    finally:
        os.environ.pop("AUTOPILOT_LLM", None)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--port", type=int, default=8765)
    ap.add_argument("--no-browser", action="store_true")
    args = ap.parse_args()
    server = ThreadingHTTPServer(("127.0.0.1", args.port), Handler)
    url = f"http://127.0.0.1:{args.port}/"
    print(f"Job Automation portal: {url}   (Ctrl+C to stop)")
    if not args.no_browser:
        threading.Timer(0.8, lambda: webbrowser.open(url)).start()
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nStopped.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
